#!/usr/bin/env python3
"""Read 공급·조합원분양·일반분양·임대 from jaegebal develop pages.

조합원분양은 재개발닷컴이 '토지등소유자 기반 추정'이라고 적는다.
시트에는 기존 열 앞에 끼우지 않고 맨 끝에 붙인다. O/P VLOOKUP 열 문자는 유지된다.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import time
from datetime import date
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent))
from apply_jaegebal_district_status import (  # noqa: E402
    GOLDEN_SHEET_ID,
    is_confident_match,
    run,
    upload_xlsx_replace,
    zone_key,
)

COUNT_RE = re.compile(r"([\d,]+)\s*(?:세대|명)")
INFO_LABELS = {
    "공급세대": "supplyUnits",
    "조합원분양": "memberUnits",
    "일반분양": "generalSaleUnits",
    "임대세대": "rentalUnits",
    "토지등소유자": "landOwners",
}
SHEET_COLUMNS = (
    "공급세대",
    "조합원분양세대",
    "일반분양세대",
    "임대세대",
    "세대수_비고",
    "세대수_확인일",
)


def first_count(text: str | None) -> int | None:
    match = COUNT_RE.search(text or "")
    if not match:
        return None
    return int(match.group(1).replace(",", ""))


def read_rows(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def scrape(page, url: str) -> dict[str, Any]:
    page.goto(url, wait_until="domcontentloaded", timeout=60000)
    page.wait_for_selector(".info-table-cell-title", timeout=20000)
    rows = page.evaluate(
        """() => [...document.querySelectorAll('.info-table-row')].map(row => ({
            title: (row.querySelector('.info-table-cell-title')?.innerText || '').trim(),
            value: (row.querySelector('.info-table-cell-value')?.innerText || '').trim()
        }))"""
    )
    found: dict[str, str] = {}
    for row in rows:
        title = re.sub(r"\s+", "", row.get("title") or "")
        for label, key in INFO_LABELS.items():
            if title == label and key not in found:
                found[key] = row.get("value") or ""
    member_text = found.get("memberUnits") or ""
    item = {
        "supplyUnits": first_count(found.get("supplyUnits")),
        "memberUnits": first_count(member_text),
        "generalSaleUnits": first_count(found.get("generalSaleUnits")),
        "rentalUnits": first_count(found.get("rentalUnits")),
        "landOwners": first_count(found.get("landOwners")),
        "memberEstimated": "추정" in member_text,
        "memberRaw": member_text,
        "pageTitle": page.title(),
    }
    parts = [item["memberUnits"], item["generalSaleUnits"], item["rentalUnits"]]
    if item["supplyUnits"] is not None and all(part is not None for part in parts):
        item["partsSum"] = sum(parts)
        item["sumMatchesSupply"] = item["partsSum"] == item["supplyUnits"]
    else:
        item["partsSum"] = None
        item["sumMatchesSupply"] = None
    return item


def launch(playwright):
    last = None
    for kwargs in ({"channel": "chrome", "headless": True}, {"headless": True}):
        try:
            return playwright.chromium.launch(**kwargs)
        except Exception as exc:  # noqa: BLE001
            last = exc
    raise SystemExit(f"Playwright launch failed ({last})")


def crawl(rows: list[dict[str, str]], delay_s: float) -> list[dict[str, Any]]:
    from playwright.sync_api import sync_playwright

    targets = []
    for row in rows:
        url = (row.get("jaegebal_url") or "").strip()
        key = zone_key(row.get("행정구"), row.get("행정동"), row.get("구역명"))
        if not url.startswith("http") or key.startswith("|"):
            continue
        targets.append((key, url, row.get("구역명") or ""))

    results: list[dict[str, Any]] = []
    with sync_playwright() as playwright:
        browser = launch(playwright)
        page = browser.new_page()
        for idx, (key, url, name) in enumerate(targets, start=1):
            item: dict[str, Any] = {"zoneNaturalKey": key, "zoneName": name, "jaegebalUrl": url}
            try:
                item.update(scrape(page, url))
                item["status"] = "ok" if item.get("supplyUnits") is not None else "missing"
            except Exception as exc:  # noqa: BLE001
                item["status"] = "error"
                item["error"] = str(exc)
            results.append(item)
            print(
                f"[{idx}/{len(targets)}] {item['status']} {key} "
                f"supply={item.get('supplyUnits')} member={item.get('memberUnits')} "
                f"general={item.get('generalSaleUnits')} rental={item.get('rentalUnits')}"
            )
            if delay_s and idx < len(targets):
                time.sleep(delay_s)
        browser.close()
    return results


def note_for(item: dict[str, Any]) -> str:
    bits = []
    if item.get("status") == "title_mismatch":
        return "페이지 제목이 구역명과 달라 세대수를 넣지 않음"
    if item.get("memberEstimated"):
        bits.append("조합원분양은 토지등소유자 기반 추정")
    if item.get("sumMatchesSupply") is False:
        bits.append(f"합계 {item.get('partsSum')} ≠ 공급세대 {item.get('supplyUnits')}")
    if item.get("supplyUnits") is not None and item.get("memberUnits") is None:
        bits.append("조합원분양·일반분양 없음")
    if item.get("status") == "error":
        bits.append("페이지 확인 실패")
    if item.get("status") == "missing":
        bits.append("공급세대 없음")
    return ". ".join(bits)


def apply_columns(src: Path, out: Path, results: list[dict[str, Any]], checked_on: str) -> dict[str, Any]:
    from openpyxl import load_workbook

    by_key = {item["zoneNaturalKey"]: item for item in results}
    wb = load_workbook(src)
    ws = wb.worksheets[0]
    headers = [cell.value for cell in ws[1]]
    hm = {str(h): i + 1 for i, h in enumerate(headers) if h}
    start = ws.max_column + 1
    for offset, name in enumerate(SHEET_COLUMNS):
        if name not in hm:
            col = start + offset
            ws.cell(1, col, name)
            hm[name] = col
    written = 0
    for r in range(2, ws.max_row + 1):
        key = zone_key(
            ws.cell(r, hm["행정구"]).value,
            ws.cell(r, hm["행정동"]).value,
            ws.cell(r, hm["구역명"]).value,
        )
        item = by_key.get(key)
        if not item:
            continue
        if item.get("pageTitle") and not is_confident_match(str(ws.cell(r, hm["구역명"]).value or ""), item["pageTitle"]):
            item["status"] = "title_mismatch"
        if item.get("status") == "ok":
            ws.cell(r, hm["공급세대"], item.get("supplyUnits"))
            ws.cell(r, hm["조합원분양세대"], item.get("memberUnits"))
            ws.cell(r, hm["일반분양세대"], item.get("generalSaleUnits"))
            ws.cell(r, hm["임대세대"], item.get("rentalUnits"))
        ws.cell(r, hm["세대수_비고"], note_for(item) or None)
        ws.cell(r, hm["세대수_확인일"], checked_on)
        written += 1
    wb.save(out)
    return {"written": written, "columns": list(SHEET_COLUMNS), "output": str(out)}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--delay", type=float, default=0.35)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--write", action="store_true")
    parser.add_argument("--sheet-id", default=GOLDEN_SHEET_ID)
    parser.add_argument("--remote", default="gdrive")
    parser.add_argument("--workdir", type=Path, default=Path("data/tmp_household_sheet"))
    args = parser.parse_args()

    if args.output.exists() and not args.apply:
        print("output exists; pass --apply to skip a new crawl")
    rows = read_rows(args.input)
    if args.apply and args.output.exists():
        payload = json.loads(args.output.read_text(encoding="utf-8"))
        results = payload["zones"]
    else:
        results = crawl(rows, args.delay)
        payload = {
            "generatedAt": date.today().isoformat(),
            "source": args.input.as_posix(),
            "zones": results,
            "ok": sum(1 for item in results if item.get("status") == "ok"),
            "missing": sum(1 for item in results if item.get("status") == "missing"),
            "error": sum(1 for item in results if item.get("status") == "error"),
            "sumMismatch": sum(1 for item in results if item.get("sumMatchesSupply") is False),
            "memberEstimated": sum(1 for item in results if item.get("memberEstimated")),
        }
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps({k: payload[k] for k in ("ok", "missing", "error", "sumMismatch", "memberEstimated")}, ensure_ascii=False))

    if not args.apply and not args.write:
        return 0

    workdir = args.workdir
    workdir.mkdir(parents=True, exist_ok=True)
    run(
        [
            "rclone",
            "backend",
            "copyid",
            f"{args.remote}:",
            args.sheet_id,
            f"{workdir}/",
            "--drive-export-formats",
            "xlsx",
        ]
    )
    files = sorted(workdir.glob("*.xlsx"))
    src = next(path for path in files if path.name != "households.xlsx")
    out = workdir / "households.xlsx"
    stats = apply_columns(src, out, results, payload["generatedAt"])
    print(json.dumps({**stats, "write": bool(args.write)}, ensure_ascii=False))
    if args.write:
        upload_xlsx_replace(args.sheet_id, out, remote=args.remote)
        print(f"https://docs.google.com/spreadsheets/d/{args.sheet_id}/edit")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
